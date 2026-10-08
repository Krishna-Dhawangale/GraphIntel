import io
import json

import docx
import fitz

from app.providers.parsers.csv_parser import CSVParser
from app.providers.parsers.docx_parser import DOCXParser
from app.providers.parsers.factory import get_parser
from app.providers.parsers.html_parser import HTMLParser
from app.providers.parsers.json_parser import JSONParser
from app.providers.parsers.pdf_parser import PDFParser
from app.providers.parsers.txt_parser import TXTParser


def test_txt_parser():
    parser = TXTParser()
    content = b"GraphIntel Phase 1\n\nFoundation Architecture and Pipeline.\n\nMarket intelligence automation."
    sections = parser.parse(content, "test.txt")
    assert len(sections) >= 1
    assert "GraphIntel" in sections[0].text


def test_html_parser():
    parser = HTMLParser()
    html_content = b"""
    <html>
        <head><title>Market Analysis Q3</title></head>
        <body>
            <h1>Executive Summary</h1>
            <p>Global revenue grew by 18% year-over-year.</p>
            <script>alert('bad');</script>
        </body>
    </html>
    """
    sections = parser.parse(html_content, "report.html")
    assert len(sections) >= 1
    full_text = " ".join([s.text for s in sections])
    assert "Executive Summary" in full_text
    assert "18% year-over-year" in full_text
    assert "alert" not in full_text


def test_csv_parser():
    parser = CSVParser()
    csv_content = b"Ticker,Revenue_Millions,Growth\nAAPL,89500,8%\nMSFT,62000,16%\nNVDA,26000,260%"
    sections = parser.parse(csv_content, "market_data.csv")
    assert len(sections) >= 2
    assert "Columns" in sections[0].text
    assert "AAPL" in sections[1].text


def test_json_parser():
    parser = JSONParser()
    data = {"company": "Acme Corp", "metrics": {"revenue": 1000000, "ebitda": 250000}}
    json_bytes = json.dumps(data).encode("utf-8")
    sections = parser.parse(json_bytes, "acme.json")
    assert len(sections) >= 2
    full_text = " ".join([s.text for s in sections])
    assert "Acme Corp" in full_text
    assert "ebitda" in full_text


def test_pdf_parser():
    # Dynamically generate a valid PDF using PyMuPDF
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 72), "GraphIntel Annual Market Intelligence Report Page 1.")
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Financial Analysis and Projections Page 2.")
    pdf_bytes = doc.write()
    doc.close()

    parser = PDFParser()
    sections = parser.parse(pdf_bytes, "annual_report.pdf")
    assert len(sections) == 2
    assert sections[0].page_number == 1
    assert "Page 1" in sections[0].text
    assert sections[1].page_number == 2
    assert "Page 2" in sections[1].text


def test_docx_parser():
    # Dynamically generate a valid DOCX using python-docx
    doc = docx.Document()
    doc.add_heading("Market Intelligence Overview", level=1)
    doc.add_paragraph("GraphIntel provides comprehensive vector and graph analytics.")
    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    parser = DOCXParser()
    sections = parser.parse(docx_bytes, "overview.docx")
    assert len(sections) >= 1
    assert "Market Intelligence Overview" in sections[0].section or "GraphIntel" in sections[0].text


def test_parser_factory():
    assert isinstance(get_parser("doc.pdf"), PDFParser)
    assert isinstance(get_parser("doc.docx"), DOCXParser)
    assert isinstance(get_parser("doc.html"), HTMLParser)
    assert isinstance(get_parser("doc.txt"), TXTParser)
    assert isinstance(get_parser("doc.csv"), CSVParser)
    assert isinstance(get_parser("doc.json"), JSONParser)
