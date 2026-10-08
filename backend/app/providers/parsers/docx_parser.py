import io
from typing import List

import docx

from app.core.exceptions import UnprocessableEntityException
from app.core.logging import logger
from app.providers.parsers.base import DocumentParser, ParsedSection


class DOCXParser(DocumentParser):
    def parse(self, content: bytes, filename: str) -> List[ParsedSection]:
        try:
            doc_file = io.BytesIO(content)
            doc = docx.Document(doc_file)
            sections: List[ParsedSection] = []

            current_section = "Introduction"
            current_paragraphs: List[str] = []

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                if p.style and p.style.name.startswith("Heading"):
                    if current_paragraphs:
                        sections.append(
                            ParsedSection(
                                text="\n\n".join(current_paragraphs),
                                section=current_section,
                                metadata={"format": "docx", "heading": current_section},
                            )
                        )
                        current_paragraphs = []
                    current_section = text
                else:
                    current_paragraphs.append(text)

            # Also parse tables if present
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        table_rows.append(row_text)
                if table_rows:
                    current_paragraphs.append("\nTable:\n" + "\n".join(table_rows))

            if current_paragraphs:
                sections.append(
                    ParsedSection(
                        text="\n\n".join(current_paragraphs),
                        section=current_section,
                        metadata={"format": "docx", "heading": current_section},
                    )
                )

            return sections
        except Exception as e:
            logger.error(f"Failed to parse DOCX {filename}: {e}")
            raise UnprocessableEntityException(f"Failed to parse DOCX document: {filename}")
